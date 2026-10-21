"""Reading the analysts' workbooks without trusting them, and writing back without breaking them."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from godavari_integration import synth_sources as ss
from godavari_integration.tracker import (WRITE_BACK_HEADER, WorkbookOpen, lock_file, migrate, parse_date,
                                          parse_file_no, read_workbook, write_back)


@pytest.fixture(scope="module")
def workbooks(tmp_path_factory, facts):
    return ss.build_trackers(tmp_path_factory.mktemp("trackers"), facts)


@pytest.fixture(scope="module")
def reads(workbooks):
    return [read_workbook(p) for p in workbooks]


def test_the_header_is_found_wherever_it_is(reads):
    assert [r.header_row for r in reads] == [3, 1, 2]
    assert all("file_number" in r.columns and "waiting_on" in r.columns for r in reads)


def test_the_total_row_is_skipped_and_the_hidden_row_is_kept_but_marked(reads):
    a = reads[0]
    assert any("total row" in s for s in a.skipped)
    hidden = [r for r in a.rows if r.hidden]
    assert [r.fields["file_number"] for r in hidden] == ["GF-24871"] and "info: hidden row" in hidden[0].problems


def test_cells_are_parsed_defensively():
    assert parse_file_no("gf-25040 ").value == "GF-25040"
    assert parse_file_no("GF-2503").value is None
    assert parse_date(46338).value.isoformat() == "2026-11-12"          # an Excel serial number
    assert parse_date("05/11/2026").value.isoformat() == "2026-11-05"   # day first: en-IN
    assert parse_date("12-Nov").value is None                           # no year and no hint: refuse to guess
    assert parse_date("11/25/2026").problem.startswith("not a day-first date")


def test_migration_sends_the_unreadable_to_a_person_and_never_guesses(reads):
    m = migrate(reads)
    assert m.rows_read == 40 and len(m.items) == 37
    assert m.needs_a_person == ["Tracker_GF-AN-1163_Nov26.xlsx:4: not a file number: 'GF-2503'"]
    assert any(c.startswith("GF-25001:") for c in m.conflicts) and m.items["GF-25001"].state == "conflict"
    assert any(x.startswith("GF-25115:") for x in m.merged_duplicates)


def test_colour_and_free_text_become_states_and_documents(reads):
    items = migrate(reads).items
    assert items["GF-25060"].state == "waiting_on_branch" and items["GF-25060"].waiting_on == ["bank_statement"]
    assert any("WhatsApp" in n for n in items["GF-25060"].notes)
    assert items["GF-25013"].state == "query_with_senior"
    assert items["GF-25035"].waiting_on == ["proof_of_address"]
    assert any("ask the analyst" in n for n in items["GF-25057"].notes)       # green, but still waiting
    assert items["GF-25129"].received_on.isoformat() == "2026-11-12"          # year from the workbook name


def test_formula_cells_without_a_cached_value_are_noted_not_trusted(reads):
    c = reads[2]
    assert all(any("formula with no cached value" in p for p in r.problems) for r in c.rows)


def test_the_customer_name_never_reaches_the_queue(reads, facts):
    items = migrate(reads).items
    text = repr(items)
    names = [f.truth.fields["applicant_name"].title() for f in facts.values()]
    assert not [n for n in names if n in text]


def test_write_back_refuses_while_excel_has_the_workbook_open(tmp_path, facts):
    [p, *_] = ss.build_trackers(tmp_path, facts)
    lock_file(p).write_text("", encoding="utf-8")
    with pytest.raises(WorkbookOpen):
        write_back(p, {"GF-25001": "written to core"})


def test_write_back_touches_only_its_own_column(tmp_path, facts):
    [p, *_] = ss.build_trackers(tmp_path, facts)
    before = load_workbook(p).worksheets[0]
    snapshot = {(c.row, c.column): (c.value, c.fill.fgColor.rgb) for row in before.iter_rows() for c in row}
    assert write_back(p, {"GF-25001": "confirmed AP0000000002", "GF-25040": "confirmed AP0000000004"}) == 2
    after = load_workbook(p).worksheets[0]
    col = next(c.column for c in after[3] if c.value == WRITE_BACK_HEADER)
    for (r, c), (value, rgb) in snapshot.items():
        if c != col:
            assert (after.cell(r, c).value, after.cell(r, c).fill.fgColor.rgb) == (value, rgb)
    assert after.row_dimensions[17].hidden and "A1:S1" in {str(m) for m in after.merged_cells.ranges}
    written = {after.cell(r, 2).value: after.cell(r, col).value for r in range(4, after.max_row + 1)}
    assert written["GF-25001"] == "confirmed AP0000000002" and written["gf-25040 "] == "confirmed AP0000000004"
