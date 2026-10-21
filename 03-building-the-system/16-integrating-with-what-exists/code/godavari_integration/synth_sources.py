"""Synthetic stand-ins for the sources Day 7's inputs inventory lists, built from Day 7's own sample.

The 27 Kukatpally files of Day 7's two-week sample are replayed as one night's 02:00 drop (Thursday 12
November 2026, week 3). Around them: the Hyderabad team's tracker workbooks, kept the way people keep
them, and the production reporting copy as of 05:00 on Friday 13 November. Every irregularity below is
deliberate and listed, so the reader's and the reconciliation's findings can be checked against the list:

  Tracker_GF-AN-1142   title row merged above the header; 'gf-25040 ' in lower case with a space; a
                       received date as an Excel serial number and one as text; GF-25147 ticked 'keyed'
                       though production holds only its shell; a hidden October row; a total row.
  Tracker_GF-AN-1157   header names spelled differently; branch written 'KPHB'; GF-25115 pasted twice;
                       '12-Nov' with no year; GF-25037 coloured green (ready for batch) though production
                       sent it back; 'bank stmt (WA recd)'; a Telugu remark.
  Tracker_GF-AN-1163   header on row 2 and an extra 'Priority' column; 'GF-2503' (a digit lost);
                       GF-25001 also claimed by GF-AN-1142; GF-25057 green with a 'waiting on' note;
                       ageing as a formula saved without a cached value.
  reporting copy       production state at 05:00; customer name and PAN present (and denied to us);
                       CODE_MST carries an APPLSTAT code 'H' added in R26.08 that mapping 1.3.0 lacks.

All names, numbers and documents are invented.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .corespec import CODE_TABLES
from .day11 import BRANCH_CODES, FileResult, GroundTruth, LoanBundle, RULESETS, Route, StubModel, build_batch, day7_sample, process
from .sftp_batch import DropFile, synthetic_pdf

NIGHT = date(2026, 11, 12)
AS_OF_RULES = date(2026, 11, 9)                    # Day 11's run date; same routes as Day 11 reports
REFRESHED_AT = datetime(2026, 11, 13, 5, 0)
NIGHT_BRANCH = "Kukatpally"
OLDER = ["GF-25016", "GF-25035", "GF-25049", "GF-25005", "GF-25017", "GF-25037", "GF-25013", "GF-25057",
         "GF-25015", "GF-25024"]
STATE_BY_BRANCH = {"Kukatpally": "telangana", "Dilsukhnagar": "telangana", "Warangal": "telangana",
                   "Vijayawada-Benz Circle": "andhra pradesh", "Guntur": "andhra pradesh",
                   "Rajahmundry": "andhra pradesh"}
REGION = {"Kukatpally": "HYD", "Dilsukhnagar": "HYD", "Warangal": "HYD", "Vijayawada-Benz Circle": "VJA",
          "Guntur": "VJA", "Rajahmundry": "VJA"}
ANALYSTS = ["GF-AN-1142", "GF-AN-1157", "GF-AN-1163"]


@dataclass
class FileFacts:
    bundle: LoanBundle
    truth: GroundTruth
    result: FileResult

    @property
    def file_no(self) -> str:
        return self.bundle.file_id

    @property
    def branch_cd(self) -> str:
        return BRANCH_CODES[self.bundle.branch]


def hyderabad_files() -> dict[str, FileFacts]:
    """Day 7's sample through Day 11's pipeline, for the three branches the Hyderabad team serves."""
    model = StubModel()
    out: dict[str, FileFacts] = {}
    for bundle, truth in build_batch(day7_sample()):
        if bundle.branch in ("Kukatpally", "Dilsukhnagar", "Warangal"):
            out[bundle.file_id] = FileFacts(bundle, truth, process(bundle, model, RULESETS[bundle.product], AS_OF_RULES))
    return out


def night_files(facts: dict[str, FileFacts]) -> list[FileFacts]:
    return [f for f in facts.values() if f.bundle.branch == NIGHT_BRANCH]


def drop_files(night: list[FileFacts]) -> list[DropFile]:
    return [DropFile(f.file_no, f.branch_cd, synthetic_pdf(f.file_no, (p.text for p in f.bundle.pages)),
                     len(f.bundle.pages)) for f in night]


def production_status(facts: dict[str, FileFacts]) -> dict[str, str]:
    """What production holds at 05:00 on 13 November (analysts key there today; our writes go to the sandbox)."""
    status: dict[str, str] = {}
    for f in night_files(facts):
        if f.result.route in (Route.READY_FOR_APPROVAL, Route.REVIEW):
            status[f.file_no] = "S" if f.file_no in ("GF-25147", "GF-25183") else "K"
        else:
            status[f.file_no] = "Q"
    status.update({"GF-25016": "B", "GF-25035": "Q", "GF-25049": "Q", "GF-25005": "Q", "GF-25017": "B",
                   "GF-25037": "Q", "GF-25013": "K", "GF-25057": "B", "GF-25015": "Q", "GF-25024": "Q",
                   "GF-24871": "A"})
    return status


# ---- the reporting copy -------------------------------------------------------------------------------
def build_reporting_db(path: Path, facts: dict[str, FileFacts], status: dict[str, str]) -> Path:
    if path.exists():
        path.unlink()
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE APPL_MST (APPL_ID TEXT PRIMARY KEY, FILE_NO TEXT, BRANCH_CD TEXT, PROD_CD TEXT, APPL_STAT TEXT,
            KYC_STAT TEXT, LOAN_AMT INTEGER, EXT_REF TEXT, ENTRY_USR TEXT, CRT_TS TEXT, LST_UPD_TS TEXT,
            APPL_NM TEXT, APPL_DOB TEXT, PAN_NO TEXT, MOB_NO TEXT, ADDR_LN1 TEXT);
        CREATE TABLE APPL_STAT_HIST (APPL_ID TEXT, OLD_STAT TEXT, NEW_STAT TEXT, CHG_TS TEXT, CHG_USR TEXT);
        CREATE TABLE FILE_REG (FILE_NO TEXT PRIMARY KEY, BRANCH_CD TEXT, REG_TS TEXT, PDF_PATH TEXT, PAGE_CNT INTEGER);
        CREATE TABLE BRANCH_MST (BRANCH_CD TEXT PRIMARY KEY, BRANCH_NM TEXT, REGION_CD TEXT, STATE_CD TEXT, PHONE_NO TEXT);
        CREATE TABLE CODE_MST (CODE_TYP TEXT, CODE_VAL TEXT, CODE_DESC TEXT, ACTIVE_FLG TEXT);
        CREATE TABLE RPT_REFRESH_LOG (REFRESH_TS TEXT, ROWS_COPIED INTEGER);
    """)
    for i, (name, code) in enumerate(BRANCH_CODES.items()):
        st = "TS" if STATE_BY_BRANCH[name] == "telangana" else "AP"
        conn.execute("INSERT INTO BRANCH_MST VALUES (?, ?, ?, ?, ?)", (code, name.upper(), REGION[name], st,
                                                                       f"040-2{i}31{i}0{i}4"))
    for typ, codes in CODE_TABLES.items():
        for val, desc in codes.items():
            conn.execute("INSERT INTO CODE_MST VALUES (?, ?, ?, 'Y')", (typ, val, desc))
    conn.execute("INSERT INTO CODE_MST VALUES ('APPLSTAT', 'H', 'ON HOLD - FRAUD CHECK', 'Y')")    # R26.08
    seq = 7_310_000
    for fno in sorted(status):
        f = facts.get(fno)
        branch_cd = f.branch_cd if f else "KKP01"
        reg = (datetime(2026, 11, 12, 2, 0) if f and f.bundle.branch == NIGHT_BRANCH
               else datetime(2026, 10, 9, 2, 0) if fno == "GF-24871"
               else datetime(2026, 11, 2 + int(fno[-1]) % 7, 2, 0))
        conn.execute("INSERT INTO FILE_REG VALUES (?, ?, ?, ?, ?)",
                     (fno, branch_cd, reg.strftime("%Y%m%d%H%M%S"), f"/region/HYD/{reg:%Y%m%d}/{fno}.pdf",
                      len(f.bundle.pages) if f else 14))
        if status[fno] == "S":
            continue
        seq += 1
        appl_id = f"AP{seq:010d}"
        keyed = reg + timedelta(hours=8, minutes=17 + seq % 40)
        g = f.truth.fields if f else {"applicant_name": "MASKED", "date_of_birth": "01-01-1980", "pan": "AAAPA0000A",
                                       "mobile": "9000000000", "address": "1-1, KPHB Colony", "requested_amount": "80000"}
        prod = {"two-wheeler": "TW", "gold": "GL"}.get(f.bundle.product if f else "", "TW")
        conn.execute("INSERT INTO APPL_MST VALUES (?, ?, ?, ?, ?, ?, ?, '', ?, ?, ?, ?, ?, ?, ?, ?)",
                     (appl_id, fno, branch_cd, prod, status[fno], "D" if status[fno] == "Q" else "C",
                      int(g["requested_amount"]), ANALYSTS[seq % 3], keyed.strftime("%Y%m%d%H%M%S"),
                      (keyed + timedelta(hours=3)).strftime("%Y%m%d%H%M%S"), g["applicant_name"],
                      g["date_of_birth"], g["pan"], g["mobile"], g["address"][:40]))
        conn.execute("INSERT INTO APPL_STAT_HIST VALUES (?, 'S', 'K', ?, ?)", (appl_id, keyed.strftime("%Y%m%d%H%M%S"), ANALYSTS[seq % 3]))
        if status[fno] != "K":
            conn.execute("INSERT INTO APPL_STAT_HIST VALUES (?, 'K', ?, ?, ?)",
                         (appl_id, status[fno], (keyed + timedelta(hours=3)).strftime("%Y%m%d%H%M%S"), ANALYSTS[seq % 3]))
    rows = conn.execute("SELECT COUNT(*) FROM APPL_MST").fetchone()[0]
    conn.execute("INSERT INTO RPT_REFRESH_LOG VALUES (?, ?)", ("20261112050000", rows - 20))
    conn.execute("INSERT INTO RPT_REFRESH_LOG VALUES (?, ?)", (REFRESHED_AT.strftime("%Y%m%d%H%M%S"), rows))
    conn.commit()
    conn.close()
    return path


# ---- the tracker workbooks -------------------------------------------------------------------------------
RED, YELLOW, GREEN = "FFFF0000", "FFFFFF00", "FF92D050"
STANDARD = ["S.No", "File No", "Branch", "Customer Name", "Product", "Loan Amt", "Assigned", "Received", "Keyed",
            "Docs checked", "Waiting on", "Verified", "Batch", "Memo", "Remarks", "Branch Ph", "Senior chk",
            "Ageing (days)", "Last upd"]
VARIANT = ["File No.", "Br", "Cust Name", "Prod", "Amount", "Analyst", "Recd", "Keyed", "Docs Chkd", "Pending with",
           "Verfied", "Batch", "Memo", "Comments", "Br Ph", "Senior Check", "Age (days)", "Last Updated", "Priority"]


def _fill(cell: object, rgb: str | None) -> None:
    if rgb:
        cell.fill = PatternFill(fill_type="solid", fgColor=rgb)  # type: ignore[attr-defined]


def _row_for(f: FileFacts | None, fno: str, analyst: str) -> dict[str, object]:
    g = f.truth.fields if f else {"applicant_name": "K SATYAVATHI", "requested_amount": "75000"}
    product = {"two-wheeler": "TW", "gold": "Gold"}.get(f.bundle.product if f else "", "TW")
    return {"file": fno, "branch": f.bundle.branch if f else "Kukatpally", "name": g["applicant_name"].title(),
            "product": product, "amount": int(g["requested_amount"]), "analyst": analyst}


def build_trackers(folder: Path, facts: dict[str, FileFacts]) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    groups = {   # Srinivas's assignment on Thursday morning: nine Kukatpally files each
        "GF-AN-1142": ["GF-25001", "GF-25040", "GF-25058", "GF-25079", "GF-25100", "GF-25122", "GF-25147", "GF-25175", "GF-25183"],
        "GF-AN-1157": ["GF-25003", "GF-25047", "GF-25060", "GF-25073", "GF-25080", "GF-25115", "GF-25129", "GF-25148", "GF-25177"],
        "GF-AN-1163": ["GF-25010", "GF-25031", "GF-25052", "GF-25075", "GF-25098", "GF-25121", "GF-25139", "GF-25169", "GF-25179"],
    }
    assert sorted(sum(groups.values(), [])) == sorted(f.file_no for f in night_files(facts))
    received_night = datetime(2026, 11, 12, 9, 5)
    paths = []

    # ---- GF-AN-1142: title row, header on row 3 ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Nov"
    ws["A1"] = "CREDIT OPS HYDERABAD - GF-AN-1142 - NOVEMBER 2026"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:S1")
    ws.append([])
    ws.append(STANDARD)
    rows = groups["GF-AN-1142"] + ["GF-25016", "GF-25035", "GF-25049", "GF-25005"]
    for n, fno in enumerate(rows, 1):
        r = _row_for(facts.get(fno), fno, "GF-AN-1142")
        shown_no = "gf-25040 " if fno == "GF-25040" else fno
        received: object = received_night
        if fno == "GF-25100":
            received = (received_night.date() - date(1899, 12, 30)).days        # typed as a serial number
        if fno == "GF-25122":
            received = "12/11/2026"
        older = fno in OLDER
        if older:
            received = datetime(2026, 11, 3 + n % 5, 10, 30)
        keyed = "Y" if fno == "GF-25147" else ("Y" if older else "")
        waiting = {"GF-25035": "addr proof since 10/11", "GF-25049": "salary slip Oct", "GF-25005": "pan copy"}.get(fno, "")
        batch = "Y" if fno == "GF-25016" else ""
        ws.append([n, shown_no, r["branch"], r["name"], r["product"], r["amount"], r["analyst"], received, keyed,
                   "Y" if older else "", waiting, "", batch, "", "", "040-23456789", "", 1 if not older else 7,
                   "12/11/2026"])
        if waiting:
            _fill(ws.cell(ws.max_row, 2), RED)
        if batch:
            _fill(ws.cell(ws.max_row, 2), GREEN)
    ws.append([len(rows) + 1, "GF-24871", "Kukatpally", "R Sujatha", "TW", 64000, "GF-AN-1142",
               datetime(2026, 10, 9), "Y", "Y", "", "Y", "Y", "Y", "sanctioned 16/10", "", "Y", 34, "16/10/2026"])
    ws.row_dimensions[ws.max_row].hidden = True
    ws.append([None, "Total files", None, None, None, None, None, None, None, None, None, None, None, None, None,
               None, None, None, None])
    ws.cell(ws.max_row, 3, f"=COUNTA(B4:B{ws.max_row - 1})")
    p = folder / "Tracker_GF-AN-1142_Nov26.xlsx"
    wb.save(p)
    paths.append(p)

    # ---- GF-AN-1157: variant header names on row 1 ----
    wb = Workbook()
    ws = wb.active
    ws.append(VARIANT)
    rows = groups["GF-AN-1157"] + ["GF-25017", "GF-25037", "GF-25013"]
    for fno in rows:
        r = _row_for(facts.get(fno), fno, "GF-AN-1157")
        older = fno in OLDER
        received: object = datetime(2026, 11, 12, 9, 40)
        if fno == "GF-25129":
            received = "12-Nov"
        if fno == "GF-25017":
            received = "05/11/2026"
        if fno in ("GF-25037", "GF-25013"):
            received = datetime(2026, 11, 6, 11, 0)
        waiting = "bank stmt (WA recd)" if fno == "GF-25060" else ""
        remark = "బ్రాంచ్ కి ఫోన్ చేశాను" if fno == "GF-25080" else ("query to senior: income" if fno == "GF-25013" else "")
        values = [fno, "KPHB" if r["branch"] == "Kukatpally" else r["branch"], r["name"], r["product"], r["amount"],
                  r["analyst"], received, "Y" if older else "", "Y" if older else "", waiting,
                  "Y" if fno in ("GF-25037", "GF-25017") else "", "Y" if fno == "GF-25017" else "", "", remark,
                  "040-23456789", "", 1, "12/11/2026", "Normal"]
        ws.append(values)
        if waiting:
            _fill(ws.cell(ws.max_row, 1), RED)
        if fno == "GF-25037":
            _fill(ws.cell(ws.max_row, 1), GREEN)
        if fno == "GF-25013":
            _fill(ws.cell(ws.max_row, 1), YELLOW)
        if fno == "GF-25115":
            ws.append(values)                                          # pasted twice
    p = folder / "Tracker_GF-AN-1157_Nov26.xlsx"
    wb.save(p)
    paths.append(p)

    # ---- GF-AN-1163: header on row 2, an extra column, a formula ----
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Updated by GF-AN-1163, Nov 2026"
    header = STANDARD[:2] + ["Priority"] + STANDARD[2:]
    ws.append(header)
    rows = groups["GF-AN-1163"] + ["GF-25001", "GF-25057", "GF-25015", "GF-25024"]
    for n, fno in enumerate(rows, 1):
        r = _row_for(facts.get(fno), fno, "GF-AN-1163")
        older = fno in OLDER
        shown_no = "GF-2503" if fno == "GF-25031" else fno
        waiting = {"GF-25057": "addr proof", "GF-25015": "bank st Sep", "GF-25024": "quotation"}.get(fno, "")
        row_no = ws.max_row + 1
        ws.append([n, shown_no, "High" if fno == "GF-25001" else "", r["branch"], r["name"], r["product"], r["amount"],
                   r["analyst"], datetime(2026, 11, 12, 9, 50) if not older else datetime(2026, 11, 5, 12, 0),
                   "Y" if older else "", "Y" if older else "", waiting, "", "Y" if fno == "GF-25057" else "", "", "",
                   "040-23456789", "", f"=TODAY()-I{row_no}", "12/11/2026"])
        if fno == "GF-25057":
            _fill(ws.cell(ws.max_row, 2), GREEN)
        elif waiting:
            _fill(ws.cell(ws.max_row, 2), RED)
    p = folder / "Tracker_GF-AN-1163_Nov26.xlsx"
    wb.save(p)
    paths.append(p)
    return paths
