"""AI-written draft for exercise 3.1: "a client for Godavari's core, SFTP and tracker". Do not run it.

The prompt it came from was short and reasonable. The code looks tidy, has type hints and a docstring, and
would pass a quick read. It has at least twelve problems that matter at Godavari; some would show up in
the first week, some at the first month end, one in front of the RBI auditor. Find them, say which Day 6
finding or owner's sentence each one breaks, and write the fix in one line each. Section 15, exercise 3.1.
"""

from __future__ import annotations

import glob
import logging
import os
import sqlite3
import time
from datetime import datetime

import paramiko
from openpyxl import load_workbook
from zeep import Client
from zeep.exceptions import Fault
from zeep.transports import Transport

log = logging.getLogger("godavari")
WSDL_URL = "http://corelend-dev01:7001/corelend/services/LendingService?wsdl"


class GodavariCoreClient:
    def __init__(self, username: str, password: str) -> None:
        self.client = Client(WSDL_URL, transport=Transport(timeout=30))
        self.username, self.password = username, password

    def build_payload(self, app: dict) -> dict:
        return {
            "FILE_NO": app["file_number"],
            "BRANCH_CD": app["branch_code"],
            "PROD_CD": "TW" if app["product"] == "two-wheeler" else "GL",
            "APPL_NM": app["applicant_name"][:60],
            "PAN_NO": app["pan"],
            "MOB_NO": app["mobile"],
            "ADDR_LN1": app["address"][:40],
            "CUST_FLG_3": "Y" if int(app.get("monthly_income") or 0) > 25000 else "N",
            "MNTH_INC": str(app.get("monthly_income") or 0),
            "LOAN_AMT": str(app["requested_amount"]),
            "TENOR_MM": str(app["tenure_months"]),
            "ROI_PCT": "14.5",
            "KYC_STAT": "C",
        }

    def already_exists(self, file_no: str) -> bool:
        conn = sqlite3.connect("reporting.sqlite")
        rows = conn.execute(f"SELECT * FROM APPL_MST WHERE FILE_NO = '{file_no}'").fetchall()
        return len(rows) > 0

    def create(self, app: dict, retries: int = 3) -> str | None:
        payload = self.build_payload(app)
        log.info("creating application %s", payload)
        if self.already_exists(app["file_number"]):
            return None
        for attempt in range(retries):
            try:
                result = self.client.service.CreateApplication(Application=payload)
                return result["APPL_ID"]
            except Fault as fault:
                log.error("core rejected %s: %s", app["file_number"], fault.message)
                return None
            except Exception:
                time.sleep(2 ** attempt)
        return None


def pick_up_nightly_batch(host: str, user: str, password: str, local_dir: str) -> list[str]:
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    ssh.connect(host, username=user, password=password)
    sftp = ssh.open_sftp()
    picked = []
    for name in sftp.listdir("/region/HYD/" + datetime.now().strftime("%Y%m%d")):
        if name.endswith(".pdf"):
            remote = f"/region/HYD/{datetime.now():%Y%m%d}/{name}"
            sftp.get(remote, os.path.join(local_dir, name))
            sftp.remove(remote)
            picked.append(name)
    return picked


def read_tracker(path: str) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, r))
        if row.get("File No"):
            received = row.get("Received")
            if isinstance(received, str):
                row["Received"] = datetime.strptime(received, "%m/%d/%Y")
            rows.append(row)
    return rows


def write_status_to_tracker(path: str, statuses: dict[str, str]) -> None:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[1].value in statuses:
            row[10].value = statuses[row[1].value]
    wb.save(path)


def nightly_reconciliation(journal: list[dict]) -> str:
    created = [j for j in journal if j.get("appl_id")]
    return f"{len(created)} applications created, 0 duplicates. All good."


def run_all(apps: list[dict]) -> None:
    core = GodavariCoreClient(os.environ["CORE_USER"], os.environ["CORE_PASSWORD"])
    for app in apps:
        core.create(app)
    for f in glob.glob("trackers/*.xlsx"):
        write_status_to_tracker(f, {a["file_number"]: "pushed" for a in apps})
